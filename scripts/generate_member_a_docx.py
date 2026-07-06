from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "成员A报告正文.md"
OUT = ROOT / "docs" / "成员A报告正文.docx"


def set_run_font(run, east_asia: str, size_pt: float, bold: bool = False):
    run.font.name = east_asia
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size_pt)
    run.bold = bold


def set_paragraph_format(paragraph, first_line: bool = False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, "宋体", 10.5, bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_cover(doc: Document):
    for _ in range(5):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("《Web 后端开发技术》期末大作业报告")
    set_run_font(run, "黑体", 18, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("成员 A 负责部分")
    set_run_font(run, "黑体", 16, True)

    for _ in range(4):
        doc.add_paragraph("")

    meta_lines = [
        "项目名称：AI 驱动的高并发秒杀与智能电商后台系统",
        "团队成员：A（姓名/学号）、B（姓名/学号）、C（姓名/学号）",
        "撰写成员：A（姓名/学号）",
        "提交日期：2026 年 6 月",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, "宋体", 12)

    doc.add_page_break()


def add_heading1(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, "黑体", 16, True)


def add_heading2(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "黑体", 14, False)


def add_body(doc: Document, text: str):
    if not text.strip():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_format(p, first_line=True)
    run = p.add_run(text.strip())
    set_run_font(run, "宋体", 12)


def add_code(doc: Document, text: str):
    for line in text.rstrip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        set_run_font(run, "Consolas", 9.5)


def add_caption(doc: Document, text: str, kind: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "黑体", 10.5, False)


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append(lines[i].strip())
        i += 1
    parsed = []
    for row in rows:
        cells = [c.strip().strip("`") for c in row.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells):
            continue
        parsed.append(cells)
    return parsed, i


def add_table(doc: Document, rows: list[list[str]], caption: str | None):
    if not rows:
        return
    if caption:
        add_caption(doc, caption, "table")
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            set_cell_text(cell, text, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph("")


def add_image_if_exists(doc: Document, image_path_text: str):
    image_path = (ROOT / image_path_text.strip()).resolve()
    if not image_path.exists():
        add_code(doc, image_path_text)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(14.5))


def build_doc():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    add_cover(doc)

    # Extract abstract/keywords and start main body from chapter 1.
    abstract_lines = []
    keywords_lines = []
    body_start = 0
    mode = None
    for idx, line in enumerate(lines):
        if line.strip() == "# 摘要":
            mode = "abstract"
            continue
        if line.strip() == "# 关键词":
            mode = "keywords"
            continue
        if line.strip() == "# 1 绪论":
            body_start = idx
            break
        if mode == "abstract":
            if line.strip():
                abstract_lines.append(line)
        elif mode == "keywords":
            if line.strip():
                keywords_lines.append(line)

    add_heading1(doc, "摘要")
    for line in abstract_lines:
        add_body(doc, line)
    add_heading1(doc, "关键词")
    for line in keywords_lines:
        add_body(doc, line)
    doc.add_page_break()

    pending_caption: str | None = None
    in_code = False
    code_lines: list[str] = []

    i = body_start
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                code_text = "\n".join(code_lines)
                if code_text.strip().startswith("docs/images/") and "\n" not in code_text.strip():
                    add_image_if_exists(doc, code_text.strip())
                else:
                    add_code(doc, code_text)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            heading = stripped[2:].strip()
            if heading not in {"成员 A 报告正文", "摘要", "关键词"}:
                add_heading1(doc, heading)
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading2(doc, stripped[3:].strip())
            i += 1
            continue

        if stripped.startswith("表 ") or stripped.startswith("图 "):
            pending_caption = stripped
            i += 1
            continue

        if stripped.startswith("|"):
            rows, next_i = parse_table(lines, i)
            add_table(doc, rows, pending_caption)
            pending_caption = None
            i = next_i
            continue

        # Drop local guidance paragraph in generated Word.
        if stripped.startswith("本文档为成员 A"):
            i += 1
            continue

        add_body(doc, stripped)
        i += 1

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
